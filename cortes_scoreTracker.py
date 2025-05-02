import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

EXCEL_FILE = "student_scores.xlsx"

def initialize_workbook():
    """Create Excel file with headers if not exists."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Score", "Status"])
        wb.save(EXCEL_FILE)

def calculate_status(score):
    return "Pass" if score >= 50 else "Fail"

def add_or_update_student(name, score):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    updated = False
    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            row[1].value = score
            row[2].value = calculate_status(score)
            updated = True
            break

    if not updated:
        ws.append([name, score, calculate_status(score)])

    wb.save(EXCEL_FILE)

def submit_data():
    name = name_entry.get().strip()
    try:
        score = int(score_entry.get())
        if not name:
            raise ValueError("Name cannot be empty")
        add_or_update_student(name, score)
        messagebox.showinfo("Success", f"Record saved for {name}.")
        name_entry.delete(0, tk.END)
        score_entry.delete(0, tk.END)
    except ValueError:
        messagebox.showerror("Invalid Input", "Please enter a valid name and integer score.")

def display_records():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    output_text.delete("1.0", tk.END)
    output_text.insert(tk.END, "All Student Records:\n")
    output_text.insert(tk.END, "-" * 40 + "\n")
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, score, status = row
        output_text.insert(tk.END, f"Name: {name}, Score: {score}, Status: {status}\n")

initialize_workbook()

# GUI TKINTER
window = tk.Tk()
window.title("Student Score Tracker")

input_frame = tk.Frame(window, pady=10)
input_frame.pack()

tk.Label(input_frame, 
        text="Student Name:",
        font=('Helvetica', 10, 'bold')).grid(row=0, column=0, sticky="e")
name_entry = tk.Entry(input_frame, width=30)
name_entry.grid(row=0, column=1, padx=5)

tk.Label(input_frame, text="Score:",
                      font=('Helvetica', 10, 'bold')).grid(row=1, column=0, sticky="e")

score_entry = tk.Entry(input_frame, width=30)
score_entry.grid(row=1, column=1, padx=5)

submit_btn = tk.Button(input_frame, 
                       text="Submit", 
                       command=submit_data,
                       font=('Helvetica', 9, 'bold'), 
                       bg="skyblue",
                       foreground="white",
                       activebackground="lightblue")

submit_btn.grid(row=2, column=0, columnspan=2, pady=10)

display_btn = tk.Button(window, text="Display All Records",
                        command=display_records,
                        font=('Helvetica', 9, 'bold'), 
                        bg="green",
                        foreground="white",
                        activebackground="lightgreen")
display_btn.pack(pady=5)

output_text = tk.Text(window, height=15, width=50)
output_text.pack(padx=10, pady=10)

window.mainloop()